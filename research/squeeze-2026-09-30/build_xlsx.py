import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.worksheet.table import Table, TableStyleInfo

R=pd.read_csv('sweep_results.csv')
COINS=['ADA','DOGE','XRP','SOL','AVAX','ETH','LTC','HBAR','LINK','SHIB','BTC','XLM','DOT','AAVE','BCH','XTZ']
GOOD=['ADA','DOGE','XRP','SOL','AVAX','ETH','LTC','HBAR','LINK','SHIB']
F=Font(name='Arial',size=10); FB=Font(name='Arial',size=10,bold=True); FH=Font(name='Arial',size=13,bold=True)
FILL_H=PatternFill('solid',fgColor='DDEBF7'); FILL_Y=PatternFill('solid',fgColor='FFF2CC'); FILL_G=PatternFill('solid',fgColor='E2EFDA'); FILL_R=PatternFill('solid',fgColor='FCE4D6')
thin=Side(style='thin',color='BBBBBB'); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
wb=Workbook()

def style_ws(ws,widths=None):
    for row in ws.iter_rows():
        for c in row:
            if c.font==Font(): c.font=F
    if widths:
        for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w

# ---------- README ----------
ws=wb.active; ws.title='Read me'
lines=[
 ('Squeeze parameter sweep — measured on all 16 Kraken US perps',FH),
 ('Built 2026-09-30 from Coinalyze aggregated perps (.A), 4h bars, 2025-10-31 → 2026-09-30 (11 months). OI in CONTRACTS. Every number here was computed, none copied.',F),
 ('',F),
 ('Design',FB),
 ('Level = prior UTC day high / low, built from 4h closes. Break = first 4h close through the level (optionally +offset%) each day, one per side per day.',F),
 ('OI read = OI change from the bar BEFORE the break to the bar AFTER it (8h) or to the break bar itself (4h). Enter at the close of the read bar. Hold = 8/12/24/36/48h. Return net of 0.1% round trip.',F),
 ('',F),
 ('Rules',FB),
 ('A_fade   = high break + OI up ≥ thr → SHORT   (the one that works)',F),
 ('chase    = high break + OI up ≥ thr → LONG    (the crowd; control)',F),
 ('D_cont   = low break + OI down ≥ thr → SHORT  (flush continuation)',F),
 ('D_bounce = low break + OI down ≥ thr → LONG   (capitulation bounce)',F),
 ('C_long   = low break + OI up ≥ thr → LONG     (shorts crowding in)',F),
 ('',F),
 ('Columns',FB),
 ('n = trades · mean = avg net return % · win = % of trades > 0 · t = t-stat of mean vs 0 · h1/h2 = mean in first / second half of the sample (both must be > 0 for a rule to be promoted) · avg_win/avg_loss/worst in %',F),
 ('',F),
 ('Scopes',FB),
 ('POOLED_16 = all coins. POOLED_10 = ADA DOGE XRP SOL AVAX ETH LTC HBAR LINK SHIB (the coins where the fade works — chosen AFTER seeing results, so treat POOLED_16 as the honest number). Per-coin rows are the coin ticker.',F),
 ('',F),
 ('Sheets',FB),
 ('Summary — per-coin verdict and the tuned rule. A_fade grids — threshold × hold heat-maps, pooled and per coin. Other rules — the same grids for chase / D_cont / D_bounce / C_long (what NOT to trade). Lookup — type any combination and see its result. All results — the flat table (17,850 rows) the program reads. Same table as sweep_results.csv / .json.',F),
 ('',F),
 ('Promotion rule (unchanged): beats unflagged baseline, both halves > 0, ≥5 coins same arrow, after fees. Only A_fade at OI ≥ 4%, 8h read, 36h hold clears it. trading.enabled stays false.',FB),
]
for i,(t,f) in enumerate(lines,1):
    c=ws.cell(row=i,column=1,value=t); c.font=f; c.alignment=Alignment(wrap_text=True,vertical='top')
ws.column_dimensions['A'].width=140

# ---------- Summary ----------
ws=wb.create_sheet('Summary')
ws['A1']='Per-coin verdict — A_fade (high break + OI up → short), 8h OI read, no offset'; ws['A1'].font=FH
hdr=['coin','n @4%/36h','mean % @4%/36h','win % @4%/36h','h1 %','h2 %','best thr %','best hold h','best mean %','best n','best win %','best t','dose-response? (3%→4%→5% mean)','verdict']
for j,h in enumerate(hdr,1):
    c=ws.cell(row=3,column=j,value=h); c.font=FB; c.fill=FILL_H; c.border=BOX; c.alignment=Alignment(wrap_text=True)
A=R[(R.rule=='A_fade')&(R.entry_offset_pct==0)&(R.oi_read_h==8)]
row=4
for coin in ['POOLED_16','POOLED_10']+COINS:
    s=A[A.scope==coin]
    base=s[(s.oi_thr_pct==4)&(s.hold_h==36)]
    ok=s[(s.n>=8)&(s.h1>0)&(s.h2>0)]
    best=ok.sort_values('mean',ascending=False).head(1)
    m3=s[(s.oi_thr_pct==3)&(s.hold_h==36)]['mean']; m4=base['mean']; m5=s[(s.oi_thr_pct==5)&(s.hold_h==36)]['mean']
    dose=' → '.join(f'{v.iloc[0]:+.2f}' if len(v) else '—' for v in (m3,m4,m5))
    if coin.startswith('POOLED'): verdict='reference'
    elif len(base) and base['mean'].iloc[0]>1.0 and base.h1.iloc[0]>0 and base.h2.iloc[0]>0: verdict='STRONG'
    elif len(base) and base['mean'].iloc[0]>1.0: verdict='STRONG (2nd half weaker)'
    elif len(base) and base['mean'].iloc[0]>0.3: verdict='MEDIUM'
    elif len(base) and base['mean'].iloc[0]>-0.3: verdict='WEAK / no edge'
    else: verdict='OPPOSITE — do not fade'
    vals=[coin,
          int(base.n.iloc[0]) if len(base) else None, base['mean'].iloc[0] if len(base) else None, base.win.iloc[0] if len(base) else None,
          base.h1.iloc[0] if len(base) else None, base.h2.iloc[0] if len(base) else None,
          best.oi_thr_pct.iloc[0] if len(best) else None, best.hold_h.iloc[0] if len(best) else None, best['mean'].iloc[0] if len(best) else None,
          int(best.n.iloc[0]) if len(best) else None, best.win.iloc[0] if len(best) else None, best.t.iloc[0] if len(best) else None, dose, verdict]
    for j,v in enumerate(vals,1):
        c=ws.cell(row=row,column=j,value=v); c.border=BOX; c.font=F
        if isinstance(v,float): c.number_format='0.00'
    fill=FILL_G if verdict.startswith('STRONG') else FILL_Y if verdict=='MEDIUM' else FILL_R if verdict.startswith('OPPOSITE') else None
    if fill: ws.cell(row=row,column=14).fill=fill
    row+=1
ws.cell(row=row+1,column=1,value='"best" = highest mean among cells with n ≥ 8 and both halves > 0. Blank = nothing survives. Dose column reads left→right as threshold rises; a rising sequence means bigger pile-in → bigger fade.').font=F
ws.cell(row=row+2,column=1,value='Source: sweep_results.csv, computed by sweep.py from raw4/*.csv (Coinalyze). Fee assumption 0.1% round trip (user-supplied, not venue-live).').font=F
style_ws(ws,[11,10,12,12,8,8,9,9,10,8,9,8,26,24])
ws.freeze_panes='B4'

# ---------- grid helper ----------
THR=sorted(R.oi_thr_pct.unique()); HOLDS=[8,12,24,36,48]
def write_grid(ws,r0,title,sub,metric='mean',fmt='0.00',scale=True):
    ws.cell(row=r0,column=1,value=title).font=FB
    ws.cell(row=r0+1,column=1,value='OI thr % \\ hold h').font=FB; ws.cell(row=r0+1,column=1).fill=FILL_H
    thr_list=sorted(sub.oi_thr_pct.unique())
    for j,h in enumerate(HOLDS,2):
        c=ws.cell(row=r0+1,column=j,value=h); c.font=FB; c.fill=FILL_H; c.border=BOX
    for i,thr in enumerate(thr_list,2):
        c=ws.cell(row=r0+i,column=1,value=thr); c.font=FB; c.border=BOX
        for j,h in enumerate(HOLDS,2):
            v=sub[(sub.oi_thr_pct==thr)&(sub.hold_h==h)]
            c=ws.cell(row=r0+i,column=j,value=(float(v[metric].iloc[0]) if len(v) else None)); c.border=BOX; c.font=F; c.number_format=fmt
    if scale and len(thr_list):
        rng=f'B{r0+2}:{get_column_letter(len(HOLDS)+1)}{r0+1+len(thr_list)}'
        ws.conditional_formatting.add(rng,ColorScaleRule(start_type='num',start_value=-2,start_color='F8696B',mid_type='num',mid_value=0,mid_color='FFFFFF',end_type='num',end_value=2,end_color='63BE7B'))
    return r0+2+len(thr_list)

def grid_block(ws,r0,label,sub):
    # mean grid, then n grid and win grid side by side to the right
    r_end=write_grid(ws,r0,f'{label} — mean net % (green = fade paid)',sub)
    # n and win to the right
    base_col=8
    for k,(metric,fmt,ttl) in enumerate([('n','0','n trades'),('win','0','win %'),('h2','0.00','2nd-half mean %')]):
        col=base_col+k*7
        ws.cell(row=r0,column=col,value=f'{label} — {ttl}').font=FB
        ws.cell(row=r0+1,column=col,value='thr \\ hold').font=FB; ws.cell(row=r0+1,column=col).fill=FILL_H
        thr_list=sorted(sub.oi_thr_pct.unique())
        for j,h in enumerate(HOLDS,1):
            c=ws.cell(row=r0+1,column=col+j,value=h); c.font=FB; c.fill=FILL_H; c.border=BOX
        for i,thr in enumerate(thr_list,2):
            c=ws.cell(row=r0+i,column=col,value=thr); c.font=FB; c.border=BOX
            for j,h in enumerate(HOLDS,1):
                v=sub[(sub.oi_thr_pct==thr)&(sub.hold_h==h)]
                c=ws.cell(row=r0+i,column=col+j,value=(float(v[metric].iloc[0]) if len(v) else None)); c.border=BOX; c.font=F; c.number_format=fmt
        if metric=='h2' and len(thr_list):
            rng=f'{get_column_letter(col+1)}{r0+2}:{get_column_letter(col+len(HOLDS))}{r0+1+len(thr_list)}'
            ws.conditional_formatting.add(rng,ColorScaleRule(start_type='num',start_value=-2,start_color='F8696B',mid_type='num',mid_value=0,mid_color='FFFFFF',end_type='num',end_value=2,end_color='63BE7B'))
    return r_end+1

# ---------- A_fade grids ----------
ws=wb.create_sheet('A_fade grids')
ws['A1']='A_fade: high break + OI up ≥ thr → SHORT. 8h OI read, no entry offset. Rows = OI threshold %, columns = hold hours.'; ws['A1'].font=FH
r=3
A8=R[(R.rule=='A_fade')&(R.entry_offset_pct==0)&(R.oi_read_h==8)]
for scope in ['POOLED_16','POOLED_10']+COINS:
    r=grid_block(ws,r,scope,A8[A8.scope==scope])
ws['A'+str(r+1)]='Variants on the pooled set (POOLED_10): OI read window and entry offset'; ws['A'+str(r+1)].font=FH; r+=3
for rd in [4,8]:
    for off in sorted(R.entry_offset_pct.unique()):
        sub=R[(R.rule=='A_fade')&(R.scope=='POOLED_10')&(R.oi_read_h==rd)&(R.entry_offset_pct==off)]
        r=grid_block(ws,r,f'POOLED_10 · OI read {rd}h · entry offset {off:.1f}%',sub)
style_ws(ws,[16]+[8]*6+[10]+[7]*6+[10]+[7]*6+[10]+[7]*6)

# ---------- Other rules ----------
ws=wb.create_sheet('Other rules')
ws['A1']='Control and secondary rules, 8h read, no offset. These are the "do not trade" evidence (chase, C_long, D_bounce) and the coin-specific D_cont.'; ws['A1'].font=FH
r=3
for rule,desc in [('chase','high break + OI up → LONG (going with the crowd)'),('D_cont','low break + OI down → SHORT (flush continuation)'),('D_bounce','low break + OI down → LONG (capitulation bounce)'),('C_long','low break + OI up → LONG (shorts crowding in)')]:
    ws.cell(row=r,column=1,value=f'{rule}: {desc}').font=FH; r+=2
    for scope in ['POOLED_16','POOLED_10']+COINS:
        sub=R[(R.rule==rule)&(R.entry_offset_pct==0)&(R.oi_read_h==8)&(R.scope==scope)]
        if len(sub): r=grid_block(ws,r,f'{rule} · {scope}',sub)
    r+=2
style_ws(ws,[16]+[8]*6+[10]+[7]*6+[10]+[7]*6+[10]+[7]*6)

# ---------- Lookup (formula-driven) ----------
ws=wb.create_sheet('Lookup')
ws['A1']='Type inputs in the yellow cells; results pull from "All results" with INDEX/MATCH.'; ws['A1'].font=FH
labels=[('rule','A_fade'),('scope','POOLED_10'),('oi_thr_pct',4),('entry_offset_pct',0),('oi_read_h',8),('hold_h',36)]
for i,(l,v) in enumerate(labels,3):
    ws.cell(row=i,column=1,value=l).font=FB
    c=ws.cell(row=i,column=2,value=v); c.fill=FILL_Y; c.font=Font(name='Arial',size=10,color='0000FF'); c.border=BOX
ws['A10']='key'; ws['A10'].font=FB
ws['B10']='=B3&"|"&B4&"|"&TEXT(B5,"0.0")&"|"&TEXT(B6,"0.0")&"|"&B7&"|"&B8'
ws['A11']='row in All results'; ws['A11'].font=FB
ws['B11']="=IFERROR(MATCH(B10,'All results'!$A:$A,0),\"no such combination\")"
outs=['n','mean','win','t','h1','h2','avg_win','avg_loss','worst']
cols={'n':'J','mean':'K','win':'L','t':'M','h1':'N','h2':'O','avg_win':'P','avg_loss':'Q','worst':'R'}
for i,o in enumerate(outs,13):
    ws.cell(row=i,column=1,value=o).font=FB
    c=ws.cell(row=i,column=2,value=f"=IF(ISNUMBER($B$11),INDEX('All results'!${cols[o]}:${cols[o]},$B$11),\"\")"); c.border=BOX; c.font=F; c.number_format='0.00'
ws['D3']='Valid values'; ws['D3'].font=FB
ws['D4']='rule: A_fade, chase, D_cont, D_bounce, C_long'; ws['D5']='scope: POOLED_16, POOLED_10, or a ticker'; ws['D6']='oi_thr_pct: 1,2,3,4,5,6,8 (A/chase/C_long) · 0,0.5,1,2,3 (D_cont/D_bounce)'
ws['D7']='entry_offset_pct: 0, 0.2, 0.5, 1'; ws['D8']='oi_read_h: 4 or 8'; ws['D9']='hold_h: 8,12,24,36,48'
style_ws(ws,[18,14,4,70])

# ---------- All results (flat) ----------
ws=wb.create_sheet('All results')
R2=R.copy()
R2.insert(0,'key',R2.rule+'|'+R2.scope+'|'+R2.oi_thr_pct.map(lambda x:f'{x:.1f}')+'|'+R2.entry_offset_pct.map(lambda x:f'{x:.1f}')+'|'+R2.oi_read_h.astype(str)+'|'+R2.hold_h.astype(str))
cols_order=['key','rule','side','arrow','scope','oi_thr_pct','entry_offset_pct','oi_read_h','hold_h','n','mean','win','t','h1','h2','avg_win','avg_loss','worst']
R2=R2[cols_order]
for j,h in enumerate(cols_order,1):
    c=ws.cell(row=1,column=j,value=h); c.font=FB; c.fill=FILL_H
for i,rowv in enumerate(R2.itertuples(index=False),2):
    for j,v in enumerate(rowv,1):
        c=ws.cell(row=i,column=j,value=(v.item() if hasattr(v,'item') else v)); c.font=F
ws.freeze_panes='A2'; ws.auto_filter.ref=f'A1:{get_column_letter(len(cols_order))}{len(R2)+1}'
style_ws(ws,[34,9,6,6,11,10,14,9,7,6,7,6,6,7,7,8,8,7])

wb.save('squeeze_parameter_sweep.xlsx'); print('saved',len(R2),'rows')
